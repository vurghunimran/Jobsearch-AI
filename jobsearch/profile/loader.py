"""Load profile.yaml and extract plain text from your CV."""

from __future__ import annotations

import functools
import logging
from pathlib import Path

import yaml

from jobsearch.config import Settings, get_settings
from jobsearch.profile.schema import Profile

log = logging.getLogger(__name__)


class ProfileError(RuntimeError):
    pass


def load_profile(settings: Settings | None = None) -> Profile:
    settings = settings or get_settings()
    path = settings.profile_path
    if not path.exists():
        raise ProfileError(
            f"No profile at {path}. Run `jobsearch init` to create a starter file, "
            "then fill it in and drop your CV into data/cv/."
        )
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(raw, dict):
        raise ProfileError(f"{path} must contain a YAML mapping at the top level.")
    return Profile.model_validate(raw)


def cv_path(profile: Profile, settings: Settings | None = None) -> Path | None:
    settings = settings or get_settings()
    if not profile.cv_file:
        return None
    candidate = Path(profile.cv_file)
    if not candidate.is_absolute():
        candidate = settings.cv_dir / candidate
    return candidate if candidate.exists() else None


def extract_text(path: Path) -> str:
    """Best-effort plain text from PDF / DOCX / TXT / MD."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if suffix in {".txt", ".md", ".markdown", ".rst"}:
        return path.read_text(encoding="utf-8", errors="replace").strip()
    raise ProfileError(
        f"Cannot read {path.name}: supported CV formats are .pdf, .docx, .txt and .md."
    )


@functools.lru_cache(maxsize=8)
def _extract_cached(path_str: str, mtime: float, size: int) -> str:
    del mtime, size  # part of the cache key only
    return extract_text(Path(path_str))


def cv_text(profile: Profile, settings: Settings | None = None) -> str:
    """CV text, re-extracted automatically whenever the file changes on disk."""
    path = cv_path(profile, settings)
    if path is None:
        return ""
    stat = path.stat()
    try:
        return _extract_cached(str(path), stat.st_mtime, stat.st_size)
    except Exception as exc:  # a corrupt PDF must not take the whole run down
        log.warning("Could not extract text from %s: %s", path, exc)
        return ""


def _work_authorization(auth) -> str:
    """State where the candidate may work, and where they may not.

    Work rights are country-specific: a Dutch permit is not an EU-wide one. A
    bare "no sponsorship required" would let a writer imply the candidate can
    take a job in any country, which is exactly the kind of false claim that
    gets an application binned.
    """
    bits: list[str] = []
    if auth.authorized_countries:
        countries = ", ".join(auth.authorized_countries)
        bits.append(f"may work without sponsorship in {countries}")
        bits.append(
            "would need sponsorship anywhere else"
            if not auth.requires_sponsorship
            else "requires visa sponsorship, including in the countries listed"
        )
    else:
        bits.append(
            "requires visa sponsorship" if auth.requires_sponsorship else "no sponsorship required"
        )
    if auth.current_visa_status:
        bits.append(f"current status: {auth.current_visa_status}")
    if auth.notice_period:
        bits.append(f"notice period {auth.notice_period}")
    if auth.willing_to_relocate:
        bits.append(f"willing to relocate ({', '.join(auth.relocation_targets) or 'open'})")
    return "; ".join(bits)


def _period(start: str, end: str) -> str:
    """Render a date range.

    An open end means the role is current. Writing just the start date reads as
    a single point in time, which invites a model to describe a multi-year job
    as though it lasted a month.
    """
    if start and end:
        return f"{start} – {end}"
    if start:
        return f"{start} – present"
    return end or "n/a"


def profile_brief(profile: Profile) -> str:
    """A compact, prompt-friendly rendering of the structured profile."""
    ident = profile.identity
    lines: list[str] = []
    if ident.full_name:
        lines.append(f"Name: {ident.full_name}")
    where = ", ".join(x for x in (ident.city, ident.country) if x)
    if where:
        lines.append(f"Based in: {where}")
    if profile.headline:
        lines.append(f"Headline: {profile.headline}")
    if profile.summary:
        lines.append(f"Summary: {profile.summary}")
    if profile.skills:
        lines.append(f"Skills: {', '.join(profile.skills)}")
    if profile.languages:
        lines.append(f"Languages: {', '.join(profile.languages)}")
    if profile.certifications:
        lines.append(f"Certifications: {', '.join(profile.certifications)}")

    lines.append("Work authorization: " + _work_authorization(profile.work_authorization))

    prefs = profile.preferences
    if prefs.min_salary:
        lines.append(
            f"Compensation expectation: at least {prefs.min_salary:,} "
            f"{prefs.salary_currency} per {prefs.salary_period}."
        )

    if profile.experience:
        lines.append("\nExperience:")
        for item in profile.experience:
            period = _period(item.start, item.end)
            head = f"- {item.title} at {item.company}"
            if item.location:
                head += f" ({item.location})"
            lines.append(f"{head}, {period}")
            if item.summary:
                lines.append(f"  {item.summary}")
            for highlight in item.highlights:
                lines.append(f"  * {highlight}")
            if item.technologies:
                lines.append(f"  tech: {', '.join(item.technologies)}")

    if profile.education:
        lines.append("\nEducation:")
        for edu in profile.education:
            period = _period(edu.start, edu.end)
            # An unstated degree type is common; don't emit a dangling "in".
            qualification = " in ".join(x for x in (edu.degree, edu.field_of_study) if x)
            line = f"- {qualification or 'Studies'}, {edu.institution}"
            if edu.location:
                line += f" ({edu.location})"
            lines.append(f"{line}, {period}")
            for highlight in edu.highlights:
                lines.append(f"  * {highlight}")

    return "\n".join(lines).strip()


STARTER_PROFILE = """\
# ---------------------------------------------------------------------------
# jobsearch-ai profile
#
# This file is the agent's only source of truth about you. Fill it in once,
# adjust the `preferences` block whenever your search changes, and put your CV
# in data/cv/ (PDF, DOCX, TXT or MD) with its filename under `cv_file`.
#
# This file stays on your machine. It is git-ignored by default.
# ---------------------------------------------------------------------------

identity:
  full_name: ""
  email: ""
  phone: ""            # include country code, e.g. "+372 5555 0000"
  city: ""
  country: ""
  linkedin: ""
  github: ""
  portfolio: ""

work_authorization:
  authorized_countries: []      # e.g. ["Estonia", "European Union"]
  requires_sponsorship: false
  current_visa_status: ""       # e.g. "EU Blue Card", "Student residence permit"
  willing_to_relocate: false
  relocation_targets: []
  notice_period: ""             # e.g. "1 month"
  earliest_start_date: ""       # e.g. "2026-10-01"

headline: ""                    # one line, e.g. "Data analyst, 3 yrs, SQL + Python"
summary: ""                     # 2-4 sentences in your own words

# The CV is attached to applications AND used as source material for writing.
cv_file: ""                     # filename inside data/cv/, e.g. "cv.pdf"
extra_documents: {}             # e.g. {transcript: "transcript.pdf"}

experience:
  - title: ""
    company: ""
    location: ""
    start: ""                   # "2023-01"
    end: ""                     # "" means present
    summary: ""
    highlights:
      - ""                      # quantified results work best
    technologies: []

education:
  - degree: ""
    field_of_study: ""
    institution: ""
    location: ""
    start: ""
    end: ""
    grade: ""
    highlights: []

skills: []
certifications: []
languages: []                   # e.g. ["English (C1)", "Azerbaijani (native)"]
links: {}

# ---------------------------------------------------------------------------
# What you are looking for. This drives the whole search.
# ---------------------------------------------------------------------------
preferences:
  titles:                       # target roles, matched loosely
    - ""
  exclude_titles: []            # e.g. ["Senior", "Manager"] if you want IC/junior only
  keywords: []                  # optional: must appear somewhere in the posting
  exclude_keywords: []          # hard veto, e.g. ["security clearance", "unpaid"]

  locations: []                 # e.g. ["Tallinn", "Estonia", "Berlin"]
  remote_types: ["remote", "hybrid", "onsite"]
  accept_worldwide_remote: true

  employment_types: ["full_time"]   # full_time | part_time | contract | internship
  seniority: []                     # intern | junior | mid | senior | lead | manager
  min_salary: null              # advisory — given to the scorer, not a hard filter
  salary_period: "year"         # "year" or "month"
  salary_currency: "EUR"
  exclude_companies: []
  max_posting_age_days: 30

# ---------------------------------------------------------------------------
# Where to search. Company boards are highest signal — add the ones you care
# about. Aggregators give breadth. `jobsearch doctor` checks each one.
# ---------------------------------------------------------------------------
sources:
  greenhouse_boards: []         # board token from job-boards.greenhouse.io/<token>
  lever_companies: []           # slug from jobs.lever.co/<slug>
  ashby_orgs: []                # org from jobs.ashbyhq.com/<org>
  workable_accounts: []         # subdomain from apply.workable.com/<sub>
  remotive: true
  arbeitnow: true
  remoteok: false
  adzuna: false                 # needs ADZUNA_APP_ID / ADZUNA_APP_KEY in .env
  adzuna_country: "gb"

# ---------------------------------------------------------------------------
# Answers to the questions nearly every form asks. Filling these in is what
# lets the agent complete an application without coming back to you.
# ---------------------------------------------------------------------------
answers:
  years_of_experience: ""
  desired_salary: ""
  available_start_date: ""
  willing_to_relocate: ""
  requires_visa_sponsorship: ""
  authorized_to_work: ""
  how_did_you_hear: "Company website"
  gender: "Decline to self-identify"
  race_ethnicity: "Decline to self-identify"
  veteran_status: "Decline to self-identify"
  disability_status: "Decline to self-identify"
  extra: {}

writing:
  tone: "warm, direct, specific"
  cover_letter_words: 280
  statement_of_purpose_words: 550
  language: "English"
  avoid_phrases:
    - "I am writing to express my interest"
    - "I believe I would be a great fit"
  notes: ""
"""


def write_starter_profile(settings: Settings | None = None, force: bool = False) -> Path:
    settings = settings or get_settings()
    settings.ensure_dirs()
    path = settings.profile_path
    if path.exists() and not force:
        raise ProfileError(f"{path} already exists. Pass --force to overwrite it.")
    path.write_text(STARTER_PROFILE, encoding="utf-8")
    return path
