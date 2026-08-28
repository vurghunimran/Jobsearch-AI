"""Render generated text to the file formats employers accept."""

from __future__ import annotations

import re
from pathlib import Path

from markdown_it import MarkdownIt

_md = MarkdownIt("commonmark", {"breaks": True, "linkify": False})


def to_html(text: str) -> str:
    """Markdown to HTML for the dashboard preview."""
    return _md.render(text or "")


def _paragraphs(text: str) -> list[str]:
    blocks = re.split(r"\n\s*\n", (text or "").strip())
    return [re.sub(r"\s*\n\s*", " ", block).strip() for block in blocks if block.strip()]


def render_pdf(text: str, path: Path, title: str = "", author: str = "") -> Path:
    """A clean single-column PDF. Body text stays selectable for ATS parsers."""
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15.5,
        alignment=TA_JUSTIFY,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "Heading",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        spaceAfter=10,
    )

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=22 * mm,
        bottomMargin=20 * mm,
        title=title or "Document",
        author=author or "",
    )
    flow = []
    if title:
        flow.append(Paragraph(_escape(title), heading))
    for block in _paragraphs(text):
        flow.append(Paragraph(_escape(block), body))
    if not flow:
        flow.append(Spacer(1, 1))
    document.build(flow)
    return path


def _escape(text: str) -> str:
    """Escape for reportlab's mini-HTML, then restore bold/italic markdown."""
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    escaped = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)
    escaped = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", escaped)
    return escaped


def render_docx(text: str, path: Path, title: str = "") -> Path:
    """A DOCX for employers whose upload widget refuses PDFs."""
    import docx
    from docx.shared import Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if title:
        document.add_heading(title, level=1)
    for block in _paragraphs(text):
        document.add_paragraph(block)
    document.save(str(path))
    return path
