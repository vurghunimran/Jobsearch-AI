from jobsearch.documents.generator import _grounding, _user_prompt, choose_document_kinds
from jobsearch.documents.render import render_docx, render_pdf, to_html
from jobsearch.matching.scorer import build_prompt as build_scoring_prompt
from jobsearch.models import Document, DocumentKind, Job


def job() -> Job:
    return Job(
        fingerprint="fp",
        source="test",
        company="ExampleCorp",
        title="Data Analyst, Growth",
        location="Tallinn, Estonia",
        remote_type="hybrid",
        employment_type="full_time",
        description="Own growth reporting. SQL, dbt.",
        url="https://example.com/job",
        rationale="Four years of exactly this work.",
        strengths=["Churn model cut churn 18%"],
        gaps=["No experimentation platform work"],
    )


class TestScoringPrompt:
    def test_carries_the_profile_cv_and_posting(self, profile):
        prompt = build_scoring_prompt(job(), profile, "Jane Rivera. dbt and SQL at Bolt.")
        assert "Jane Rivera" in prompt
        assert "dbt and SQL at Bolt" in prompt
        assert "Data Analyst, Growth" in prompt
        assert "Own growth reporting" in prompt

    def test_truncates_a_huge_description(self, profile):
        big = job()
        big.description = "word " * 20000
        prompt = build_scoring_prompt(big, profile, "cv")
        assert "[...truncated...]" in prompt
        assert len(prompt) < 40000

    def test_states_missing_pieces_rather_than_leaving_holes(self, profile):
        bare = job()
        bare.description = ""
        prompt = build_scoring_prompt(bare, profile, "")
        assert "(no description provided by the source)" in prompt
        assert "(no CV file configured)" in prompt


class TestWritingPrompt:
    def test_includes_the_earlier_assessment_so_gaps_are_handled_honestly(self, profile):
        prompt = _user_prompt(job(), profile, "cv text", "Write the cover letter.")
        assert "Four years of exactly this work." in prompt
        assert "Churn model cut churn 18%" in prompt
        assert "No experimentation platform work" in prompt
        assert "never fake" in prompt

    def test_carries_the_users_rewrite_instruction(self, profile):
        prompt = _user_prompt(job(), profile, "cv", "Write it.", note="Lead with the churn work.")
        assert "Lead with the churn work." in prompt
        assert "candidate's instructions" in prompt

    def test_screening_prompts_omit_the_fit_assessment(self, profile):
        prompt = _user_prompt(job(), profile, "cv", "Answer.", include_fit=False)
        assert "Prior assessment" not in prompt


class TestGrounding:
    def test_bans_fabrication_and_placeholders(self, profile):
        text = _grounding(profile)
        assert "Never invent" in text
        assert "No placeholders" in text

    def test_carries_the_users_style_settings(self, profile):
        profile.writing.tone = "dry and understated"
        profile.writing.language = "German"
        profile.writing.avoid_phrases = ["synergy"]
        text = _grounding(profile)
        assert "dry and understated" in text
        assert "German" in text
        assert "synergy" in text


class TestDocumentSelection:
    def test_uses_the_scorers_recommendation(self, profile):
        target = job()
        target.recommended_documents = ["statement_of_purpose"]
        assert choose_document_kinds(target, profile) == [DocumentKind.statement_of_purpose]

    def test_falls_back_to_a_cover_letter(self, profile):
        assert choose_document_kinds(job(), profile) == [DocumentKind.cover_letter]

    def test_ignores_unknown_or_unwritable_kinds(self, profile):
        target = job()
        target.recommended_documents = ["screening_answers", "nonsense"]
        assert choose_document_kinds(target, profile) == [DocumentKind.cover_letter]


class TestRendering:
    SAMPLE = "Dear team,\n\nI cut latency **40%** at Bolt.\n\nBest,\nJane"

    def test_pdf_is_a_real_pdf(self, tmp_path):
        path = render_pdf(self.SAMPLE, tmp_path / "letter.pdf", "Cover letter", "Jane Rivera")
        assert path.read_bytes().startswith(b"%PDF")

    def test_docx_round_trips_the_text(self, tmp_path):
        import docx

        render_docx(self.SAMPLE, tmp_path / "letter.docx", "Cover letter")
        text = "\n".join(p.text for p in docx.Document(str(tmp_path / "letter.docx")).paragraphs)
        assert "I cut latency" in text and "Jane" in text

    def test_empty_content_still_produces_a_file(self, tmp_path):
        assert render_pdf("", tmp_path / "empty.pdf").exists()

    def test_angle_brackets_do_not_break_the_pdf_renderer(self, tmp_path):
        # reportlab parses a mini-HTML dialect, so raw <> must be escaped.
        path = render_pdf("I use <script> tags & ampersands", tmp_path / "x.pdf")
        assert path.stat().st_size > 0

    def test_markdown_becomes_html_for_the_dashboard(self):
        assert "<strong>40%</strong>" in to_html(self.SAMPLE)


class TestPacketDocumentSelection:
    def test_cover_letter_text_prefers_the_letter(self, profile, seeded, settings):
        from jobsearch.apply.base import ApplicationPacket
        from jobsearch.models import Application

        documents = [
            Document(application_id=1, kind=DocumentKind.resume_summary, content="Summary."),
            Document(application_id=1, kind=DocumentKind.cover_letter, content="The letter."),
        ]
        packet = ApplicationPacket(job(), Application(job_id=1), profile, documents)
        assert packet.cover_letter_text() == "The letter."

    def test_falls_back_when_there_is_no_cover_letter(self, profile):
        from jobsearch.apply.base import ApplicationPacket
        from jobsearch.models import Application

        documents = [
            Document(application_id=1, kind=DocumentKind.statement_of_purpose, content="My SOP.")
        ]
        packet = ApplicationPacket(job(), Application(job_id=1), profile, documents)
        assert packet.cover_letter_text() == "My SOP."
