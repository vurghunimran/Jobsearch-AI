import pytest

from jobsearch.profile.loader import (
    ProfileError,
    cv_text,
    extract_text,
    load_profile,
    profile_brief,
    write_starter_profile,
)
from jobsearch.profile.schema import Identity, Profile


class TestStarterProfile:
    def test_written_file_parses_back_into_the_schema(self, settings):
        write_starter_profile(settings)
        assert load_profile(settings) is not None

    def test_refuses_to_clobber_without_force(self, settings):
        write_starter_profile(settings)
        with pytest.raises(ProfileError, match="already exists"):
            write_starter_profile(settings)
        write_starter_profile(settings, force=True)

    def test_a_missing_profile_says_what_to_run(self, settings):
        with pytest.raises(ProfileError, match="jobsearch init"):
            load_profile(settings)


class TestCvExtraction:
    def test_reads_markdown(self, settings, profile):
        assert "churn model" in cv_text(profile, settings)

    def test_re_extracts_when_the_file_changes(self, settings, profile):
        assert "churn model" in cv_text(profile, settings)
        (settings.cv_dir / "cv.md").write_text("Totally different CV.", encoding="utf-8")
        assert "Totally different" in cv_text(profile, settings)

    def test_reads_docx(self, settings, tmp_path):
        import docx

        document = docx.Document()
        document.add_paragraph("Jane Rivera")
        document.add_paragraph("Data analyst with dbt experience.")
        path = tmp_path / "cv.docx"
        document.save(str(path))
        assert "dbt experience" in extract_text(path)

    def test_reads_pdf(self, settings, tmp_path):
        from jobsearch.documents.render import render_pdf

        path = render_pdf("Jane Rivera, data analyst at Bolt.", tmp_path / "cv.pdf")
        assert "Jane Rivera" in extract_text(path)

    def test_rejects_an_unsupported_format(self, tmp_path):
        path = tmp_path / "cv.pages"
        path.write_text("x")
        with pytest.raises(ProfileError, match="supported CV formats"):
            extract_text(path)

    def test_no_cv_configured_is_not_an_error(self, settings):
        assert cv_text(Profile(), settings) == ""

    def test_a_corrupt_file_degrades_instead_of_crashing(self, settings, profile):
        (settings.cv_dir / "cv.pdf").write_bytes(b"not really a pdf")
        profile.cv_file = "cv.pdf"
        assert cv_text(profile, settings) == ""


class TestProfileBrief:
    def test_includes_the_facts_a_writer_needs(self, profile):
        brief = profile_brief(profile)
        assert "Jane Rivera" in brief
        assert "Tallinn" in brief
        assert "Work authorization" in brief

    def test_states_sponsorship_either_way(self, profile):
        assert "no sponsorship required" in profile_brief(profile)
        profile.work_authorization.requires_sponsorship = True
        assert "requires visa sponsorship" in profile_brief(profile)

    def test_empty_profile_does_not_crash(self):
        assert isinstance(profile_brief(Profile()), str)

    def test_an_ongoing_role_is_marked_present(self, profile):
        from jobsearch.profile.schema import Experience

        profile.experience = [Experience(title="Analyst", company="Acme", start="2026-03", end="")]
        brief = profile_brief(profile)
        # Bare "2026-03" reads as a point in time and invites a writer to
        # describe a current role as if it had already ended.
        assert "2026-03 – present" in brief

    def test_a_degree_without_a_named_qualification_reads_cleanly(self, profile):
        from jobsearch.profile.schema import Education

        profile.education = [
            Education(degree="", field_of_study="European Studies", institution="Maastricht")
        ]
        brief = profile_brief(profile)
        assert "European Studies, Maastricht" in brief
        assert " in European Studies" not in brief, "no dangling 'in' when degree is blank"

    def test_a_named_degree_still_reads_normally(self, profile):
        from jobsearch.profile.schema import Education

        profile.education = [
            Education(degree="BA", field_of_study="International Studies", institution="ADA")
        ]
        assert "BA in International Studies, ADA" in profile_brief(profile)


class TestIdentity:
    def test_splits_a_full_name_when_parts_are_absent(self):
        ident = Identity(full_name="Jane Maria Rivera")
        assert ident.resolved_first_name() == "Jane"
        assert ident.resolved_last_name() == "Rivera"

    def test_explicit_parts_win(self):
        ident = Identity(full_name="Ignored Name", first_name="Jane", last_name="Rivera")
        assert (ident.resolved_first_name(), ident.resolved_last_name()) == ("Jane", "Rivera")

    def test_single_word_name_has_no_surname(self):
        assert Identity(full_name="Prince").resolved_last_name() == ""
